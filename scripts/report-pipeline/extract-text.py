#!/usr/bin/env python
"""Best-effort local text extraction for report-pipeline inputs.

The script writes extracted text to stdout and exits 0 even when a file has no
extractable text. The caller decides whether empty text is release-blocking.
"""

from __future__ import annotations

import csv
import io
import pathlib
import sys


def extract_pdf(file_path: pathlib.Path) -> str:
    errors: list[str] = []

    try:
        import fitz  # PyMuPDF

        with fitz.open(file_path) as doc:
            pages = [page.get_text("text").strip() for page in doc]
        text = "\n\n".join(page for page in pages if page)
        if text.strip():
            return text
    except Exception as exc:  # pragma: no cover - environment dependent fallback
        errors.append(f"fitz: {exc}")

    try:
        import pdfplumber

        with pdfplumber.open(file_path) as pdf:
            pages = [(page.extract_text() or "").strip() for page in pdf.pages]
        text = "\n\n".join(page for page in pages if page)
        if text.strip():
            return text
    except Exception as exc:  # pragma: no cover - environment dependent fallback
        errors.append(f"pdfplumber: {exc}")

    try:
        from pypdf import PdfReader

        reader = PdfReader(str(file_path))
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
        text = "\n\n".join(page for page in pages if page)
        if text.strip():
            return text
    except Exception as exc:  # pragma: no cover - environment dependent fallback
        errors.append(f"pypdf: {exc}")

    return ""


def extract_docx(file_path: pathlib.Path) -> str:
    try:
        import docx

        document = docx.Document(str(file_path))
        return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
    except Exception:
        return ""


def extract_xlsx(file_path: pathlib.Path) -> str:
    try:
        import openpyxl

        workbook = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
        rows: list[list[str]] = []
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(value.strip() for value in values):
                    rows.append(values)
        output = []
        for row in rows:
            buffer = io.StringIO()
            writer = csv.writer(buffer)
            writer.writerow(row)
            output.append(buffer.getvalue().strip())
        return "\n".join(output)
    except Exception:
        return ""


def extract_text(file_path: pathlib.Path) -> str:
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(file_path)
    if suffix == ".docx":
        return extract_docx(file_path)
    if suffix == ".xlsx":
        return extract_xlsx(file_path)
    if suffix in {".txt", ".md", ".csv"}:
        return file_path.read_text(encoding="utf-8", errors="replace")
    return ""


def main(argv: list[str]) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")

    if len(argv) != 2:
        print("Usage: extract-text.py <file>", file=sys.stderr)
        return 2

    file_path = pathlib.Path(argv[1])
    if not file_path.exists():
        print(f"File not found: {file_path}", file=sys.stderr)
        return 2

    text = extract_text(file_path)
    sys.stdout.write(text)
    if text and not text.endswith("\n"):
        sys.stdout.write("\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
